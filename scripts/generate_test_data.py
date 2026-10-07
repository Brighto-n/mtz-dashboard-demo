import random, os, csv
from datetime import datetime, timedelta, date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

random.seed(2026)
OUT="Тестові_дані"; UNITS_DIR=os.path.join(OUT,"Файли_частин")
os.makedirs(UNITS_DIR,exist_ok=True)

# ---------- ієрархія: групи прямого підпорядкування (назви — від користувача), усе нижче вигадане ----------
GROUPS=[("G1","Група А"),("G2","Група Б"),("G3","Група В"),("G4","Група Г"),("G5","Група Д")]
TARGET=1600
cs_root=["Граніт","Сосна","Ясен","Кремінь","Бастіон","Вежа","Явір","Скеля","Рубіж","Ворон","Тополя","Сапсан",
 "Бурштин","Обрій","Каштан","Причал","Маяк","Яструб","Дуб","Ожина","Терен","Калина","Пагорб","Вітер","Кордон",
 "Лелека","Глід","Рута","Клен","Острів","Бриз","Ліхтар","Цитадель","Беркут","Вільха","Чайка","Комета","Зубр",
 "Борсук","Кремінець","Дзвін","Пролісок","Гарт","Сармат","Колос","Тайфун","Аргон","Сокіл","Стриж","Гроза"]
bat_types=["1 мб","2 мб","3 мб","4 мб","1 шб","2 шб","1 тб","рота БпАК","мінбатр","1 рр"]
surn=["Коваленко","Мельник","Шевчук","Бондар","Ткачук","Кравець","Олійник","Гнатюк","Савчук","Лисенко","Руденко",
 "Мороз","Литвин","Павлюк","Гуменюк","Остапенко","Дорошенко","Кушнір","Гаврилюк","Марченко"]
ranks=["сол.","ст. сол.","мол. серж.","серж.","ст. серж."]

units=[]; subs=[]; positions=[]
n=1001; gi=0
while len(positions)<TARGET:
    gc,gname=GROUPS[gi%len(GROUPS)]; gi+=1
    vch=f"Т{n}"; n+=1
    separate = random.random()<0.35          # окрема частина без батальйонів
    units.append((vch,gc,"окрема частина" if separate else "бригада/полк"))
    bats=["основний склад"] if separate else random.sample(bat_types,random.randint(3,5))
    for st in bats:
        sc=f"{vch}-{st}"; subs.append((sc,st,vch))
        k=min(random.randint(2,10),TARGET-len(positions))
        if k<=0: break
        for cs in random.sample(cs_root,k):
            positions.append(dict(code=f"{sc}-{cs}",name=cs,sub=st,vch=vch,group=gc,pers=random.randint(2,6)))
group_d=dict(GROUPS); unit_group={u[0]:u[1] for u in units}
for p in positions: p["path"]=f"{p['group']}/{p['vch']}/{p['sub']}/{p['name']}"

CATS=[("G","Набір продуктів, кг","food_kg"),("H","Пайок, кг","ration_kg"),("I","Вода, кг","water_kg"),
 ("J","ЗУ, кг","zu_kg"),("K","ЗББ, кг","zbb_kg"),("L","АтППО, кг","atppo_kg"),("M","речове майно, кг","clothing_kg"),
 ("N","ПММ,кг","fuel_kg"),("O","засоби звʼязку/АКБ","comms_kg"),("P","Інше МТЗ, кг","other_kg")]
METHODS={"НРК":(["НРК колісний","НРК гусеничний"],60,180),
         "БпЛА":(["дрон-доставщик важкий","дрон-доставщик легкий"],5,20),
         "пішим порядком":(["група підносу"],10,35),
         "автотранспорт":(["пікап","броньований автомобіль"],80,400)}
method_w=[("НРК",4),("БпЛА",4),("пішим порядком",2),("автотранспорт",1)]
notes=["доставлено в повному обсязі"]*6+["доставлено частково","перенесено через обстріл","зрив доставки, повтор наступного дня","АКБ на зарядку повернуто","потрібна вода додатково"]

def wchoice(lst): return random.choices([x for x,_ in lst],[w for _,w in lst])[0]
def split(total,pers):
    w={"food_kg":3*pers,"ration_kg":2*pers,"water_kg":6*pers,"zu_kg":random.choice([0,0,0,2,5]),
       "zbb_kg":random.choice([0,0,4,8,12]),"atppo_kg":random.choice([0,0,0,1,3]),"clothing_kg":random.choice([0,0,0,2,4]),
       "fuel_kg":random.choice([0,0,5,10]),"comms_kg":random.choice([0,1,2,3]),"other_kg":random.choice([0,0,1,2,4])}
    s=sum(w.values()); f=total/s
    return {k:round(v*f,1) for k,v in w.items()}

start=date(2026,9,28); days=10
records=[]; did=1
for p in positions:
    d=0
    while d<days:
        day=start+timedelta(days=d)
        m=wchoice(method_w); veh,lo,hi=METHODS[m]
        total=random.uniform(lo,hi)
        kg=split(total,p["pers"])
        note=random.choice(notes)
        if note.startswith("зрив"): kg={k:0.0 for k in kg}
        dt=datetime(day.year,day.month,day.day,random.choice([4,5,6,19,20,21,22,23]),random.choice([0,10,20,30,40,50]))
        r=ranks[random.randrange(len(ranks))]+" "+random.choice(surn)
        records.append(dict(id=did,dt=dt,pos=p,kg=kg,method=m,veh=random.choice(veh),resp=r,
             link=f"https://example.org/s/test{did:05d}",note=note,entry=dt+timedelta(hours=random.choice([0,1,2,3,6,12,26]))))
        did+=1; d+=random.choice([1,1,1,1,2])

# ---------- стилі ----------
F="Arial"; thin=Side(style="thin",color="BBBBBB"); B=Border(left=thin,right=thin,top=thin,bottom=thin)
HF=Font(name=F,bold=True,color="FFFFFF",size=10); HFILL=PatternFill("solid",fgColor="3B4A3F")
INP=PatternFill("solid",fgColor="FFF6CC"); CALC=PatternFill("solid",fgColor="E8E8E8"); base=Font(name=F,size=10)
HEAD=["№ п/п ","військова частина ","підрозділ ","Позиція ","кількість в/с на позиції ","Дата, час доставки"]+[c[1] for c in CATS]+\
     ["загальна вага, кг","відповідальний за пакування ","спосіб доставки ","спосіб доставки","посилання на пітвердження доставки","примітка","Дата запису"]
WID=[6,14,12,14,12,17]+[11]*10+[12,24,16,22,34,30,17]

def unit_file(vch):
    recs=sorted([r for r in records if r["pos"]["vch"]==vch],key=lambda r:(r["dt"],r["pos"]["sub"],r["pos"]["name"]))
    pos=[p for p in positions if p["vch"]==vch]
    wb=Workbook(); ws=wb.active; ws.title="Доставка"
    for i,(h,w) in enumerate(zip(HEAD,WID),1):
        c=ws.cell(row=1,column=i,value=h); c.font=HF; c.fill=HFILL; c.border=B
        c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width=w
    ws.row_dimensions[1].height=45; ws.freeze_panes="E2"
    MAXR=max(500,len(recs)+300)
    for r_i in range(2,MAXR+2):
        for c in range(1,24):
            cell=ws.cell(row=r_i,column=c); cell.font=base; cell.border=B
            if c==17:
                cell.value=f'=IF(COUNT(G{r_i}:P{r_i})=0,"",SUM(G{r_i}:P{r_i}))'; cell.fill=CALC
            elif c==1:
                cell.value=f'=IF(D{r_i}="","",ROW()-1)'; cell.fill=CALC
            else:
                cell.fill=INP; cell.protection=Protection(locked=False)
            if 7<=c<=17: cell.number_format="0.0"
            if c in (6,23): cell.number_format="DD.MM.YYYY HH:MM"
    for i,r in enumerate(recs,2):
        p=r["pos"]
        vals=[None,vch,p["sub"],p["name"],p["pers"],r["dt"]]+[r["kg"][c[2]] for c in CATS]+[None,r["resp"],r["method"],r["veh"],r["link"],r["note"],r["entry"]]
        for c,v in enumerate(vals,1):
            if v is not None: ws.cell(row=i,column=c,value=v)
    # довідник для списків
    dsh=wb.create_sheet("Довідник")
    dsh["A1"]="Підрозділи"; dsh["B1"]="Позиції"; dsh["C1"]="Спосіб доставки"
    for i,s in enumerate(sorted({p["sub"] for p in pos}),2): dsh.cell(row=i,column=1,value=s)
    for i,p in enumerate(sorted({p["name"] for p in pos}),2): dsh.cell(row=i,column=2,value=p)
    for i,m in enumerate(METHODS,2): dsh.cell(row=i,column=3,value=m)
    nsub=len({p["sub"] for p in pos}); npos=len({p["name"] for p in pos})
    dsh.sheet_state="hidden"
    ws.add_data_validation(dv:=DataValidation(type="list",formula1=f'"{vch}"')); dv.add(f"B2:B{MAXR+1}")
    ws.add_data_validation(dv:=DataValidation(type="list",formula1="Довідник!$A$2:$A$300")); dv.add(f"C2:C{MAXR+1}")
    ws.add_data_validation(dv:=DataValidation(type="list",formula1="Довідник!$B$2:$B$300")); dv.add(f"D2:D{MAXR+1}")
    ws.add_data_validation(dv:=DataValidation(type="list",formula1=f"Довідник!$C$2:$C${len(METHODS)+1}")); dv.add(f"S2:S{MAXR+1}")
    ws.add_data_validation(dv:=DataValidation(type="whole",operator="between",formula1="0",formula2="50",showErrorMessage=True,error="Ціле число 0–50")); dv.add(f"E2:E{MAXR+1}")
    ws.add_data_validation(dv:=DataValidation(type="decimal",operator="between",formula1="0",formula2="2000",showErrorMessage=True,error="Вага в кг: число 0–2000")); dv.add(f"G2:P{MAXR+1}")
    ws.add_data_validation(dv:=DataValidation(type="date",operator="greaterThan",formula1="45658",showErrorMessage=True,error="Введіть дату й час")); dv.add(f"F2:F{MAXR+1}"); dv.add(f"W2:W{MAXR+1}")
    ws.protection.sheet=True; ws.protection.autoFilter=False; ws.protection.sort=False
    ws.auto_filter.ref=f"A1:W{MAXR+1}"
    d=os.path.join(UNITS_DIR,group_d[unit_group[vch]]); os.makedirs(d,exist_ok=True)
    fn=os.path.join(d,f"{vch}_доставка.xlsx"); wb.save(fn); return fn

files=[unit_file(u[0]) for u in units]

# ---------- довідник ієрархії ----------
hier=[]
for c,nm in GROUPS: hier.append((c,nm,"група прямого підпорядкування","",c))
for v,gc,kind in units: hier.append((v,f"в/ч {v}",kind,gc,f"{gc}/{v}"))
for sc,st,v in subs: hier.append((sc,st,"підрозділ",v,f"{unit_group[v]}/{v}/{st}"))
for p in positions: hier.append((p["code"],p["name"],"позиція",f"{p['vch']}-{p['sub']}",p["path"]))
with open(os.path.join(OUT,"dovidnyk_pidrozdiliv.csv"),"w",newline="",encoding="utf-8-sig") as f:
    w=csv.writer(f); w.writerow(["code","name","level","parent_code","path"]); w.writerows(hier)
os.makedirs(os.path.join(OUT,"Зведення"),exist_ok=True)
with open(os.path.join(OUT,"Зведення","stan_zapysiv.csv"),"w",newline="",encoding="utf-8-sig") as f:
    w=csv.writer(f); w.writerow(["key","first_seen"])
    for r in records:
        p=r["pos"]; ed=r["dt"]+timedelta(hours=0)
        w.writerow(["|".join([p["vch"].lower(),p["sub"].lower(),p["name"].lower(),r["dt"].strftime("%Y-%m-%d %H:%M")]),r["entry"].date().isoformat()])
print(len(units),"units",len(subs),"subs",len(positions),"positions",len(records),"records")
