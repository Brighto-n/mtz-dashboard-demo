"""
Оновлення файлів частин за довідником (для розширення).

Як додати позицію/підрозділ/частину:
  1. Додайте рядок(и) в dovidnyk_pidrozdiliv.csv (code,name,level,parent_code,path).
     path: Група/в_ч/підрозділ/позиція, напр. G2/Т1100/1 мб/Нова
  2. Запустіть:  python onovyty_shablony.py dovidnyk_pidrozdiliv.csv Файли_частин
Що робить:
  - для нової частини створює порожній файл <Група>/<в_ч>_доставка.xlsx;
  - в наявних файлах оновлює випадні списки підрозділів і позицій (внесені дані не чіпає).
Потрібно: Python 3 + openpyxl.
"""
import sys, os, csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HEAD=["№ п/п ","військова частина ","підрозділ ","Позиція ","кількість в/с на позиції ","Дата, час доставки",
 "Набір продуктів, кг","Пайок, кг","Вода, кг","ЗУ, кг","ЗББ, кг","АтППО, кг","речове майно, кг","ПММ,кг",
 "засоби звʼязку/АКБ","Інше МТЗ, кг","загальна вага, кг","відповідальний за пакування ","спосіб доставки ",
 "спосіб доставки","посилання на пітвердження доставки","примітка","Дата запису"]
WID=[6,14,12,14,12,17]+[11]*10+[12,24,16,22,34,30,17]
METHODS=["НРК","БпЛА","пішим порядком","автотранспорт"]
F="Arial"; thin=Side(style="thin",color="BBBBBB"); B=Border(left=thin,right=thin,top=thin,bottom=thin)

def new_file(fn,vch,rows=1000):
    wb=Workbook(); ws=wb.active; ws.title="Доставка"
    for i,(h,w) in enumerate(zip(HEAD,WID),1):
        c=ws.cell(row=1,column=i,value=h); c.font=Font(name=F,bold=True,color="FFFFFF",size=10)
        c.fill=PatternFill("solid",fgColor="3B4A3F"); c.border=B
        c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width=w
    ws.row_dimensions[1].height=45; ws.freeze_panes="E2"
    inp=PatternFill("solid",fgColor="FFF6CC"); calc=PatternFill("solid",fgColor="E8E8E8")
    for r in range(2,rows+2):
        for c in range(1,24):
            cell=ws.cell(row=r,column=c); cell.font=Font(name=F,size=10); cell.border=B
            if c==1: cell.value=f'=IF(D{r}="","",ROW()-1)'; cell.fill=calc
            elif c==17: cell.value=f'=IF(COUNT(G{r}:P{r})=0,"",SUM(G{r}:P{r}))'; cell.fill=calc
            else: cell.fill=inp; cell.protection=Protection(locked=False)
            if 7<=c<=17: cell.number_format="0.0"
            if c in (6,23): cell.number_format="DD.MM.YYYY HH:MM"
    d=wb.create_sheet("Довідник"); d.sheet_state="hidden"
    L=rows+1
    for rng,f in [(f"B2:B{L}",f'"{vch}"'),(f"C2:C{L}","Довідник!$A$2:$A$300"),(f"D2:D{L}","Довідник!$B$2:$B$300"),(f"S2:S{L}","Довідник!$C$2:$C$20")]:
        v=DataValidation(type="list",formula1=f); ws.add_data_validation(v); v.add(rng)
    for rng,kw in [(f"E2:E{L}",dict(type="whole",operator="between",formula1="0",formula2="50",error="Ціле число 0–50")),
                   (f"G2:P{L}",dict(type="decimal",operator="between",formula1="0",formula2="2000",error="Вага в кг: число 0–2000")),
                   (f"F2:F{L}",dict(type="date",operator="greaterThan",formula1="45658",error="Введіть дату й час"))]:
        v=DataValidation(showErrorMessage=True,**kw); ws.add_data_validation(v); v.add(rng)
        if rng.startswith("F"): v.add(f"W2:W{L}")
    ws.protection.sheet=True; ws.protection.autoFilter=False; ws.protection.sort=False
    ws.auto_filter.ref=f"A1:W{L}"
    wb.save(fn)

def refresh_lists(fn,subs,poss):
    wb=load_workbook(fn); d=wb["Довідник"] if "Довідник" in wb.sheetnames else wb.create_sheet("Довідник")
    for col in "ABC":
        for r in range(2,301): d[f"{col}{r}"]=None
    d["A1"],d["B1"],d["C1"]="Підрозділи","Позиції","Спосіб доставки"
    for i,v in enumerate(sorted(subs),2): d.cell(row=i,column=1,value=v)
    for i,v in enumerate(sorted(poss),2): d.cell(row=i,column=2,value=v)
    for i,v in enumerate(METHODS,2): d.cell(row=i,column=3,value=v)
    d.sheet_state="hidden"; wb.save(fn)

def main(dov,root):
    rows=list(csv.DictReader(open(dov,encoding="utf-8-sig")))
    names={r["code"]:r["name"] for r in rows}
    units={}
    for r in rows:
        p=r["path"].split("/")
        if len(p)>=2: units.setdefault(p[1],{"group":names.get(p[0],p[0]),"subs":set(),"pos":set()})
        if len(p)>=3: units[p[1]]["subs"].add(p[2])
        if len(p)==4: units[p[1]]["pos"].add(p[3])
    created=updated=0
    for vch,u in sorted(units.items()):
        d=os.path.join(root,u["group"]); os.makedirs(d,exist_ok=True)
        fn=os.path.join(d,f"{vch}_доставка.xlsx")
        if not os.path.exists(fn): new_file(fn,vch); created+=1
        else: updated+=1
        refresh_lists(fn,u["subs"],u["pos"])
    print(f"Створено нових файлів: {created}, оновлено списків: {updated}")

if __name__=="__main__":
    if len(sys.argv)!=3: print(__doc__); sys.exit(1)
    main(*sys.argv[1:])
